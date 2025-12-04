# # """
# # Plagiarism Detection and Removal Module
# # Contains functionality for detecting and removing plagiarism from text
# # """

# # import re
# # import anthropic
# # import logging
# # from typing import Dict, List, Tuple
# # from difflib import SequenceMatcher

# # logger = logging.getLogger(__name__)

# # class PlagiarismDetector:
# #     def __init__(self, anthropic_api_key: str, ai_detector=None):
# #         # Initialize Anthropic client for plagiarism removal
# #         self.anthropic_available = anthropic_api_key != "dummy_key"
# #         if self.anthropic_available:
# #             try:
# #                 self.client = anthropic.Anthropic(api_key=anthropic_api_key)
# #                 logger.info("✅ Anthropic client initialized for plagiarism removal")
# #             except Exception as e:
# #                 logger.warning(f"⚠️ Anthropic client failed: {e}")
# #                 self.client = None
# #                 self.anthropic_available = False
# #         else:
# #             self.client = None
# #             logger.warning("⚠️ Anthropic client disabled for plagiarism removal")

# #         # Reference to AI detector for pattern analysis
# #         self.ai_detector = ai_detector
        
# #         logger.info("✅ Plagiarism Detector initialized")

# #     def check_plagiarism_local(self, text1: str, text2: str) -> Tuple[float, dict]:
# #         """Check for plagiarism between two texts using local similarity analysis"""
# #         if not text1.strip() or not text2.strip():
# #             return 0.0, {"error": "Both texts must be provided"}

# #         # Clean and normalize texts
# #         text1_clean = re.sub(r'[^\w\s]', '', text1.lower())
# #         text2_clean = re.sub(r'[^\w\s]', '', text2.lower())

# #         # Overall similarity using SequenceMatcher
# #         overall_similarity = SequenceMatcher(None, text1_clean, text2_clean).ratio()

# #         # Sentence-by-sentence comparison
# #         sentences1 = [s.strip() for s in re.split(r'[.!?]+', text1) if s.strip()]
# #         sentences2 = [s.strip() for s in re.split(r'[.!?]+', text2) if s.strip()]

# #         sentence_similarities = []
# #         for s1 in sentences1:
# #             max_sim = 0.0
# #             for s2 in sentences2:
# #                 sim = SequenceMatcher(None, s1.lower(), s2.lower()).ratio()
# #                 if sim > max_sim:
# #                     max_sim = sim
# #             sentence_similarities.append(max_sim)

# #         avg_sentence_similarity = sum(sentence_similarities) / len(sentence_similarities) if sentence_similarities else 0.0

# #         # Word overlap analysis
# #         words1 = set(text1_clean.split())
# #         words2 = set(text2_clean.split())
# #         if words1 and words2:
# #             word_overlap = len(words1.intersection(words2)) / len(words1.union(words2))
# #         else:
# #             word_overlap = 0.0

# #         # N-gram analysis (3-grams)
# #         def get_ngrams(text, n=3):
# #             words = text.split()
# #             return set(zip(*[words[i:] for i in range(n)])) if len(words) >= n else set()

# #         ngrams1 = get_ngrams(text1_clean)
# #         ngrams2 = get_ngrams(text2_clean)
# #         if ngrams1 and ngrams2:
# #             ngram_overlap = len(ngrams1.intersection(ngrams2)) / len(ngrams1.union(ngrams2))
# #         else:
# #             ngram_overlap = 0.0

# #         # Combined plagiarism score
# #         plagiarism_score = (
# #             (overall_similarity * 0.4) +
# #             (avg_sentence_similarity * 0.3) +
# #             (word_overlap * 0.2) +
# #             (ngram_overlap * 0.1)
# #         )

# #         details = {
# #             'overall_similarity': overall_similarity * 100,
# #             'sentence_similarity': avg_sentence_similarity * 100,
# #             'word_overlap': word_overlap * 100,
# #             'ngram_overlap': ngram_overlap * 100,
# #             'word_count_1': len(text1.split()),
# #             'word_count_2': len(text2.split())
# #         }

# #         return plagiarism_score * 100, details

# #     def get_plagiarism_report(self, text1: str, text2: str) -> Tuple[str, float, str]:
# #         """Generate comprehensive plagiarism analysis report"""
# #         if not text1.strip() or not text2.strip():
# #             return "❌ **Please provide both texts for plagiarism comparison.**", 0, "❌ Insufficient data"

# #         plagiarism_score, details = self.check_plagiarism_local(text1, text2)

# #         # Determine classification
# #         if plagiarism_score >= 80:
# #             classification = "🚨 **HIGHLY LIKELY PLAGIARIZED**"
# #             confidence = "Very High"
# #         elif plagiarism_score >= 60:
# #             classification = "⚠️ **POSSIBLY PLAGIARIZED**"
# #             confidence = "High"
# #         elif plagiarism_score >= 40:
# #             classification = "🤔 **SOME SIMILARITIES DETECTED**"
# #             confidence = "Medium"
# #         elif plagiarism_score >= 20:
# #             classification = "📝 **MINOR SIMILARITIES**"
# #             confidence = "Low"
# #         else:
# #             classification = "✅ **LIKELY ORIGINAL**"
# #             confidence = "Very Low Plagiarism Risk"

# #         report = f"{classification}\n\n"
# #         report += f"**Plagiarism Score: {plagiarism_score:.1f}%** (Confidence: {confidence})\n\n"

# #         # Detailed analysis
# #         report += "**📊 SIMILARITY ANALYSIS:**\n\n"
# #         report += f"• **Overall Text Similarity:** {details['overall_similarity']:.1f}%\n"
# #         report += f"• **Sentence-level Similarity:** {details['sentence_similarity']:.1f}%\n"
# #         report += f"• **Word Overlap:** {details['word_overlap']:.1f}%\n"
# #         report += f"• **Phrase Pattern Overlap:** {details['ngram_overlap']:.1f}%\n\n"

# #         # Text statistics
# #         report += "**📏 TEXT STATISTICS:**\n\n"
# #         report += f"• **Text 1 Word Count:** {details['word_count_1']}\n"
# #         report += f"• **Text 2 Word Count:** {details['word_count_2']}\n"

# #         length_diff = abs(details['word_count_1'] - details['word_count_2'])
# #         report += f"• **Length Difference:** {length_diff} words\n\n"

# #         # Recommendations
# #         report += "**💡 RECOMMENDATIONS:**\n\n"
# #         if plagiarism_score >= 60:
# #             report += "• ⚠️ High similarity detected - review for potential plagiarism\n"
# #             report += "• 🔍 Check for proper citations and attribution\n"
# #             report += "• ✏️ Consider paraphrasing or rewriting similar sections\n"
# #         elif plagiarism_score >= 40:
# #             report += "• 🤔 Moderate similarities found - verify originality\n"
# #             report += "• 📝 Review overlapping sections for proper attribution\n"
# #         elif plagiarism_score >= 20:
# #             report += "• ✅ Minor similarities are normal for related topics\n"
# #             report += "• 📚 Ensure any quotes or references are properly cited\n"
# #         else:
# #             report += "• ✅ Texts appear to be original with minimal overlap\n"
# #             report += "• 🎯 Good originality indicators detected\n"

# #         status = f"Analysis complete: {classification}"

# #         return report, plagiarism_score, status

# #     def remove_plagiarism(self, text: str, rewrite_mode: str = "balanced", reference_text: str = "") -> dict:
# #         """Advanced plagiarism removal with rewriting"""
# #         if not self.anthropic_available:
# #             return {"error": "Plagiarism removal requires ANTHROPIC_API_KEY"}

# #         if not text.strip():
# #             return {"error": "Text cannot be empty"}

# #         try:
# #             # Original plagiarism score if reference provided
# #             original_plagiarism_score = 0.0
# #             if reference_text.strip():
# #                 original_plagiarism_score, _ = self.check_plagiarism_local(text, reference_text)

# #             # Get AI patterns to avoid (if AI detector is available)
# #             detected_ai_patterns = {}
# #             flagged_phrases = []
            
# #             if self.ai_detector:
# #                 detected_ai_patterns = self.ai_detector.detect_ai_patterns(text)
# #                 for patterns in detected_ai_patterns.values():
# #                     flagged_phrases.extend(patterns)

# #             # Define rewriting strategies
# #             strategies = {
# #                 "conservative": {
# #                     "description": "Light paraphrasing while maintaining original structure",
# #                     "prompt": """Rewrite this text to eliminate plagiarism while keeping the original meaning and structure as much as possible. Make minimal changes but ensure the text is sufficiently different from the original.

# # Guidelines:
# # - Replace key terms with synonyms
# # - Slightly modify sentence structure
# # - Keep the same overall organization
# # - Maintain technical accuracy
# # - Preserve the original tone
# # - Avoid AI writing patterns and clichés"""
# #                 },
# #                 "balanced": {
# #                     "description": "Moderate restructuring with synonym replacement",
# #                     "prompt": """Rewrite this text to completely eliminate plagiarism while maintaining the core meaning and information. Use a balanced approach that changes both vocabulary and sentence structure.

# # Guidelines:
# # - Replace words with appropriate synonyms
# # - Restructure sentences significantly
# # - Change paragraph organization if needed
# # - Use different transitions and connectors
# # - Maintain factual accuracy
# # - Keep the same level of detail
# # - Write in a natural, human style
# # - Avoid corporate buzzwords and AI patterns"""
# #                 },
# #                 "aggressive": {
# #                     "description": "Complete restructuring with new writing style",
# #                     "prompt": """Completely rewrite this text from scratch to eliminate all plagiarism. Transform the writing style, structure, and approach while preserving all the essential information and meaning.

# # Guidelines:
# # - Use completely different sentence structures
# # - Reorganize information flow
# # - Change the writing style and tone
# # - Use different examples or explanations where possible
# # - Replace all possible words with synonyms
# # - Create a fresh perspective on the same content
# # - Write naturally like a human would
# # - Avoid all AI writing patterns, corporate jargon, and clichés"""
# #                 }
# #             }

# #             strategy = strategies.get(rewrite_mode, strategies["balanced"])

# #             # Create the rewriting prompt
# #             full_prompt = f"""{strategy['prompt']}

# # Text to rewrite:
# # \"\"\"
# # {text}
# # \"\"\"

# # IMPORTANT INSTRUCTIONS:
# # - Write naturally and conversationally
# # - Avoid these AI patterns: {', '.join(flagged_phrases[:10]) if flagged_phrases else 'corporate buzzwords and clichés'}
# # - Use varied sentence lengths and structures
# # - Make it sound human-written

# # Return only the rewritten version with no explanations."""

# #             # Add reference context if provided
# #             if reference_text.strip():
# #                 full_prompt += f"\n\nCRITICAL: Ensure the rewritten text is significantly different from this reference text:\n\"\"\"\n{reference_text[:500]}{'...' if len(reference_text) > 500 else ''}\n\"\"\""

# #             response = self.client.messages.create(
# #                 model="claude-3-5-sonnet-20241022",
# #                 max_tokens=4000,
# #                 messages=[{"role": "user", "content": full_prompt}]
# #             )

# #             rewritten_text = response.content[0].text.strip()

# #             # Calculate improvement metrics
# #             new_plagiarism_score = 0.0
# #             if reference_text.strip():
# #                 new_plagiarism_score, _ = self.check_plagiarism_local(rewritten_text, reference_text)

# #             # Calculate AI score improvement (if AI detector is available)
# #             original_ai_score = 0.0
# #             new_ai_score = 0.0
            
# #             if self.ai_detector:
# #                 original_ai_score, _ = self.ai_detector.calculate_ai_score(text)
# #                 new_ai_score, _ = self.ai_detector.calculate_ai_score(rewritten_text)

# #             # Calculate text statistics
# #             original_word_count = len(text.split())
# #             new_word_count = len(rewritten_text.split())
# #             length_change = ((new_word_count - original_word_count) / original_word_count) * 100 if original_word_count > 0 else 0

# #             return {
# #                 "rewritten_text": rewritten_text,
# #                 "original_plagiarism_score": original_plagiarism_score,
# #                 "new_plagiarism_score": new_plagiarism_score,
# #                 "original_ai_score": original_ai_score,
# #                 "new_ai_score": new_ai_score,
# #                 "improvement": original_plagiarism_score - new_plagiarism_score,
# #                 "ai_improvement": original_ai_score - new_ai_score,
# #                 "original_word_count": original_word_count,
# #                 "new_word_count": new_word_count,
# #                 "length_change": length_change,
# #                 "rewrite_mode": rewrite_mode,
# #                 "strategy_description": strategy["description"]
# #             }

# #         except Exception as e:
# #             logger.error(f"Plagiarism removal error: {str(e)}")
# #             return {"error": str(e)}

# #     def find_similar_sentences(self, text1: str, text2: str, threshold: float = 0.7) -> List[dict]:
# #         """Find sentences with high similarity between two texts"""
# #         sentences1 = [s.strip() for s in re.split(r'[.!?]+', text1) if s.strip()]
# #         sentences2 = [s.strip() for s in re.split(r'[.!?]+', text2) if s.strip()]
        
# #         similar_pairs = []
        
# #         for i, s1 in enumerate(sentences1):
# #             for j, s2 in enumerate(sentences2):
# #                 similarity = SequenceMatcher(None, s1.lower(), s2.lower()).ratio()
# #                 if similarity >= threshold:
# #                     similar_pairs.append({
# #                         'text1_sentence': s1,
# #                         'text2_sentence': s2,
# #                         'similarity': similarity * 100,
# #                         'text1_position': i,
# #                         'text2_position': j
# #                     })
        
# #         # Sort by similarity descending
# #         similar_pairs.sort(key=lambda x: x['similarity'], reverse=True)
        
# #         return similar_pairs

# #     def get_plagiarism_statistics(self, text1: str, text2: str) -> dict:
# #         """Get detailed plagiarism statistics"""
# #         plagiarism_score, details = self.check_plagiarism_local(text1, text2)
# #         similar_sentences = self.find_similar_sentences(text1, text2)
        
# #         return {
# #             'overall_score': plagiarism_score,
# #             'details': details,
# #             'similar_sentences_count': len(similar_sentences),
# #             'similar_sentences': similar_sentences[:5],  # Top 5 most similar
# #             'total_sentences_text1': len([s for s in re.split(r'[.!?]+', text1) if s.strip()]),
# #             'total_sentences_text2': len([s for s in re.split(r'[.!?]+', text2) if s.strip()])
# #         }



# """
# plagiarism_detector.py

# Plagiarism Detection & Removal:
# - Uses Anthropic Claude if ANTHROPIC_API_KEY is set.
# - Falls back to rule-based rewrites using ai_patterns.json and phrasal_patterns.xlsx.
# - Safe, robust, and compatible with existing main.py (expects .remove()).
# """

# from __future__ import annotations
# import os
# import re
# import json
# import time
# import logging
# import random
# from pathlib import Path
# from typing import Dict, Tuple, List
# from difflib import SequenceMatcher

# # optional import (openpyxl might not be installed in some environments)
# try:
#     import openpyxl
# except Exception:
#     openpyxl = None

# # optional anthopic client import
# try:
#     import anthropic
# except Exception:
#     anthropic = None

# # logging
# logger = logging.getLogger(__name__)
# if not logger.handlers:
#     logging.basicConfig(level=logging.INFO)


# # ---------------- Helpers to load rule files ---------------- #
# BASE_DIR = Path(__file__).resolve().parent


# def load_json_patterns(filename: str = "ai_patterns.json") -> Dict[str, str]:
#     path = BASE_DIR / filename
#     if not path.exists():
#         logger.info(f"JSON patterns file not found at {path}")
#         return {}
#     try:
#         with path.open("r", encoding="utf-8") as f:
#             data = json.load(f)
#             # ensure mapping str->str and lowercase keys for safer matching
#             return {str(k).strip().lower(): str(v).strip() for k, v in data.items()}
#     except Exception as e:
#         logger.warning(f"Failed to load JSON patterns ({path}): {e}")
#         return {}


# def load_xlsx_patterns(filename: str = "phrasal_patterns.xlsx") -> Dict[str, str]:
#     path = BASE_DIR / filename
#     if not path.exists():
#         logger.info(f"XLSX patterns file not found at {path}")
#         return {}
#     if openpyxl is None:
#         logger.warning("openpyxl not installed; skipping XLSX patterns")
#         return {}
#     patterns = {}
#     try:
#         wb = openpyxl.load_workbook(path, read_only=True)
#         sheet = wb.active
#         # Expect two columns: original -> replacement. Skip header row if present.
#         for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
#             if i == 1:
#                 # skip header if header looks like strings "original" / "replacement"
#                 header = [str(c).lower() if c else "" for c in row]
#                 if "original" in header or "replacement" in header:
#                     continue
#             if not row:
#                 continue
#             orig = row[0]
#             repl = row[1] if len(row) > 1 else None
#             if orig and repl:
#                 patterns[str(orig).strip().lower()] = str(repl).strip()
#         logger.info(f"Loaded {len(patterns)} XLSX patterns from {path}")
#     except Exception as e:
#         logger.warning(f"Failed to load XLSX patterns ({path}): {e}")
#     return patterns


# # ---------------- PlagiarismDetector class ---------------- #
# class PlagiarismDetector:
#     def __init__(self, anthropic_api_key: str | None = None, ai_detector=None):
#         """
#         anthropic_api_key: pass in or it will read from ANTHROPIC_API_KEY env var.
#         ai_detector: optional object used for AI pattern detection / scoring (safe if None).
#         """
#         self.ai_detector = ai_detector

#         # Load key from env if not passed
#         self.api_key = anthropic_api_key or os.getenv("ANTHROPIC_API_KEY")
#         self.anthropic_available = bool(self.api_key) and (anthropic is not None)

#         # Load rule-based patterns (fallback)
#         self.json_patterns = load_json_patterns("ai_patterns.json")
#         self.xlsx_patterns = load_xlsx_patterns("phrasal_patterns.xlsx")

#         # Initialize Anthropic client if available
#         self.client = None
#         if self.anthropic_available:
#             try:
#                 self.client = anthropic.Anthropic(api_key=self.api_key)
#                 logger.info("✅ Anthropic client initialized for plagiarism removal")
#             except Exception as e:
#                 logger.warning(f"⚠️ Could not initialize Anthropic client: {e}")
#                 self.client = None
#                 self.anthropic_available = False
#         else:
#             if not self.api_key:
#                 logger.info("Anthropic API key not provided; using fallback rewrite patterns")
#             else:
#                 logger.info("Anthropic client library not installed; using fallback rewrite patterns")

#         logger.info("PlagiarismDetector initialized")

#     # ---------------- Local similarity checks ---------------- #
#     def check_plagiarism_local(self, text1: str, text2: str) -> Tuple[float, dict]:
#         """Return score (0-100) and details dict"""
#         if not text1 or not text2:
#             return 0.0, {"error": "Both texts must be provided"}

#         # normalize
#         t1 = re.sub(r"[^\w\s]", "", text1.lower())
#         t2 = re.sub(r"[^\w\s]", "", text2.lower())

#         overall_similarity = SequenceMatcher(None, t1, t2).ratio()

#         sentences1 = [s.strip() for s in re.split(r"[.!?]+", text1) if s.strip()]
#         sentences2 = [s.strip() for s in re.split(r"[.!?]+", text2) if s.strip()]

#         sentence_similarities = []
#         for s1 in sentences1:
#             max_sim = 0.0
#             for s2 in sentences2:
#                 sim = SequenceMatcher(None, s1.lower(), s2.lower()).ratio()
#                 if sim > max_sim:
#                     max_sim = sim
#             sentence_similarities.append(max_sim)
#         avg_sentence_similarity = sum(sentence_similarities) / len(sentence_similarities) if sentence_similarities else 0.0

#         words1 = set(t1.split())
#         words2 = set(t2.split())
#         word_overlap = len(words1.intersection(words2)) / len(words1.union(words2)) if words1 and words2 else 0.0

#         def ngrams(text: str, n: int = 3):
#             words = text.split()
#             return set(zip(*[words[i:] for i in range(n)])) if len(words) >= n else set()

#         ngrams1 = ngrams(t1)
#         ngrams2 = ngrams(t2)
#         ngram_overlap = len(ngrams1.intersection(ngrams2)) / len(ngrams1.union(ngrams2)) if ngrams1 and ngrams2 else 0.0

#         score = (overall_similarity * 0.4 + avg_sentence_similarity * 0.3 + word_overlap * 0.2 + ngram_overlap * 0.1) * 100

#         details = {
#             "overall_similarity": overall_similarity * 100,
#             "sentence_similarity": avg_sentence_similarity * 100,
#             "word_overlap": word_overlap * 100,
#             "ngram_overlap": ngram_overlap * 100,
#             "word_count_1": len(text1.split()),
#             "word_count_2": len(text2.split()),
#         }
#         return score, details

#     def get_plagiarism_report(self, text1: str, text2: str) -> Tuple[str, float, str]:
#         report, score, status = "", 0.0, "Unknown"
#         try:
#             score, details = self.check_plagiarism_local(text1, text2)
#             # classification
#             if score >= 80:
#                 classification = "🚨 HIGHLY LIKELY PLAGIARIZED"
#                 conf = "Very High"
#             elif score >= 60:
#                 classification = "⚠️ POSSIBLY PLAGIARIZED"
#                 conf = "High"
#             elif score >= 40:
#                 classification = "🤔 SOME SIMILARITIES DETECTED"
#                 conf = "Medium"
#             elif score >= 20:
#                 classification = "📝 MINOR SIMILARITIES"
#                 conf = "Low"
#             else:
#                 classification = "✅ LIKELY ORIGINAL"
#                 conf = "Very Low"

#             report = f"{classification}\n\nPlagiarism Score: {score:.1f}% (Confidence: {conf})\n\n"
#             report += "**SIMILARITY ANALYSIS**\n"
#             report += f"Overall: {details['overall_similarity']:.1f}%\n"
#             report += f"Sentence-level: {details['sentence_similarity']:.1f}%\n"
#             report += f"Word overlap: {details['word_overlap']:.1f}%\n"
#             report += f"Phrase overlap: {details['ngram_overlap']:.1f}%\n"
#             status = f"Analysis complete: {classification}"
#             return report, score, status
#         except Exception as e:
#             logger.error(f"Error building plagiarism report: {e}")
#             return "Error generating report", 0.0, "Error"

#     # ---------------- Anthropic helper (retry) ---------------- #
#     def _call_anthropic_with_retry(self, prompt: str, model: str = "claude-3-5-sonnet-latest", max_retries: int = 3) -> str:
#         if not self.client:
#             raise RuntimeError("Anthropic client not initialized")
#         last_exc = None
#         for attempt in range(1, max_retries + 1):
#             try:
#                 resp = self.client.messages.create(
#                     model=model,
#                     max_tokens=1200,
#                     messages=[{"role": "user", "content": prompt}]
#                 )
#                 # response content access is client dependent; handle safely
#                 if hasattr(resp, "content") and resp.content:
#                     # many Anthropic clients return list-like content
#                     candidate = resp.content[0].text if hasattr(resp.content[0], "text") else str(resp.content[0])
#                     return candidate.strip()
#                 # fallback to str(resp)
#                 return str(resp).strip()
#             except Exception as e:
#                 last_exc = e
#                 s = str(e).lower()
#                 if "overloaded" in s or "529" in s:
#                     if attempt < max_retries:
#                         wait = (2 ** (attempt - 1)) + random.uniform(0, 1)
#                         logger.info(f"Anthropic overloaded; retrying in {wait:.1f}s (attempt {attempt})")
#                         time.sleep(wait)
#                         continue
#                     logger.error("Anthropic overloaded and retries exhausted")
#                     raise
#                 # non-retryable: re-raise
#                 logger.warning(f"Anthropic call failed (attempt {attempt}): {e}")
#                 raise
#         # if loop finishes unexpectedly
#         raise last_exc or RuntimeError("Anthropic call failed")

#     # ---------------- Rule-based fallback rewrite ---------------- #
#     def _apply_rules(self, text: str) -> str:
#         new_text = text

#         # Apply JSON patterns (word/phrase -> replacement)
#         for k, v in self.json_patterns.items():
#             # word boundary replace ignoring case
#             try:
#                 new_text = re.sub(rf"\b{re.escape(k)}\b", v, new_text, flags=re.IGNORECASE)
#             except re.error:
#                 # skip malformed regex keys
#                 new_text = new_text.replace(k, v)

#         # Apply XLSX patterns
#         for k, v in self.xlsx_patterns.items():
#             try:
#                 new_text = re.sub(rf"\b{re.escape(k)}\b", v, new_text, flags=re.IGNORECASE)
#             except re.error:
#                 new_text = new_text.replace(k, v)

#         return new_text

#     # ---------------- Public remove() ---------------- #
#     def remove(self, text: str, rewrite_mode: str = "balanced", reference_text: str = "") -> Tuple[str, Dict]:
#         """
#         Main entry used by main.py:
#         returns (rewritten_text, report_dict)
#         report_dict contains metadata: status, counts, scores (if computed), fallback_used bool.
#         """
#         if not text or not text.strip():
#             return text, {"status": "empty_text", "rewritten_text": "", "fallback_used": False}

#         report = {
#             "status": "unknown",
#             "fallback_used": False,
#             "original_word_count": len(text.split()),
#             "new_word_count": None,
#             "original_plagiarism_score": None,
#             "new_plagiarism_score": None,
#         }

#         # Attempt Anthropic first if available
#         if self.client:
#             try:
#                 # build a clear prompt using rewrite_mode and optional reference_text guidance
#                 prompt = f"Rewrite the following text to remove plagiarism while preserving the meaning.\nRewrite mode: {rewrite_mode}\n\nText:\n{text}\n\nReturn only the rewritten text, no explanations."
#                 if reference_text:
#                     prompt += f"\n\nReference (ensure dissimilar to this):\n{reference_text[:1000]}"

#                 rewritten = self._call_anthropic_with_retry(prompt)
#                 report.update({
#                     "status": "anthropic_success",
#                     "fallback_used": False,
#                     "new_word_count": len(rewritten.split())
#                 })

#                 # if reference provided, compute simple local plagiarism improvement
#                 if reference_text:
#                     orig_score, _ = self.check_plagiarism_local(text, reference_text)
#                     new_score, _ = self.check_plagiarism_local(rewritten, reference_text)
#                     report["original_plagiarism_score"] = orig_score
#                     report["new_plagiarism_score"] = new_score

#                 return rewritten, report

#             except Exception as e:
#                 logger.warning(f"Anthropic rewrite failed, falling back to rule-based: {e}")

#         # Fallback rewrite (rules)
#         try:
#             rewritten = self._apply_rules(text)
#             report.update({
#                 "status": "fallback_success",
#                 "fallback_used": True,
#                 "new_word_count": len(rewritten.split())
#             })

#             if reference_text:
#                 orig_score, _ = self.check_plagiarism_local(text, reference_text)
#                 new_score, _ = self.check_plagiarism_local(rewritten, reference_text)
#                 report["original_plagiarism_score"] = orig_score
#                 report["new_plagiarism_score"] = new_score

#             return rewritten, report

#         except Exception as e:
#             logger.error(f"Fallback rewrite failed: {e}")
#             # As a last resort, return original text with error report
#             return text, {"status": "error", "message": str(e), "fallback_used": True}

#     # alias for compatibility with older code
#     def remove_plagiarism(self, *args, **kwargs):
#         return self.remove(*args, **kwargs)

#     # ---------------- Helpful utilities ---------------- #
#     def find_similar_sentences(self, text1: str, text2: str, threshold: float = 0.7) -> List[dict]:
#         sentences1 = [s.strip() for s in re.split(r"[.!?]+", text1) if s.strip()]
#         sentences2 = [s.strip() for s in re.split(r"[.!?]+", text2) if s.strip()]
#         pairs = []
#         for i, s1 in enumerate(sentences1):
#             for j, s2 in enumerate(sentences2):
#                 sim = SequenceMatcher(None, s1.lower(), s2.lower()).ratio()
#                 if sim >= threshold:
#                     pairs.append({
#                         "text1_sentence": s1,
#                         "text2_sentence": s2,
#                         "similarity": sim * 100,
#                         "text1_position": i,
#                         "text2_position": j
#                     })
#         pairs.sort(key=lambda x: x["similarity"], reverse=True)
#         return pairs

#     def get_plagiarism_statistics(self, text1: str, text2: str) -> Dict:
#         score, details = self.check_plagiarism_local(text1, text2)
#         pairs = self.find_similar_sentences(text1, text2)
#         return {
#             "overall_score": score,
#             "details": details,
#             "similar_sentences_count": len(pairs),
#             "similar_sentences": pairs[:5],
#             "total_sentences_text1": len([s for s in re.split(r"[.!?]+", text1) if s.strip()]),
#             "total_sentences_text2": len([s for s in re.split(r"[.!?]+", text2) if s.strip()])
#         }





"""
plagiarism_module.py
Logic only. No FastAPI routes here.
"""
from __future__ import annotations
import os
import re
import json
import time
import logging
import random
from pathlib import Path
from typing import Dict, Tuple, List
from difflib import SequenceMatcher

# Optional imports
try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    import anthropic
except Exception:
    anthropic = None

# Logging setup
logger = logging.getLogger(__name__)
if not logger.handlers:
    logging.basicConfig(level=logging.INFO)

BASE_DIR = Path(__file__).resolve().parent

# ---------------- Helpers ---------------- #

def load_json_patterns(filename: str = "ai_patterns.json") -> Dict[str, str]:
    path = BASE_DIR / filename
    if not path.exists():
        return {}
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
            return {str(k).strip().lower(): str(v).strip() for k, v in data.items()}
    except Exception:
        return {}

def load_xlsx_patterns(filename: str = "phrasal_patterns.xlsx") -> Dict[str, str]:
    path = BASE_DIR / filename
    if not path.exists() or openpyxl is None:
        return {}
    patterns = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        sheet = wb.active
        for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if i == 1: continue
            orig = row[0]
            repl = row[1] if len(row) > 1 else None
            if orig and repl:
                patterns[str(orig).strip().lower()] = str(repl).strip()
    except Exception:
        pass
    return patterns

# ---------------- Class Definition ---------------- #

class PlagiarismDetector:
    def __init__(self, anthropic_api_key: str | None = None, ai_detector=None):
        self.ai_detector = ai_detector
        self.api_key = anthropic_api_key or os.getenv("ANTHROPIC_API_KEY")
        self.anthropic_available = bool(self.api_key) and (anthropic is not None)
        
        self.json_patterns = load_json_patterns("ai_patterns.json")
        self.xlsx_patterns = load_xlsx_patterns("phrasal_patterns.xlsx")
        
        self.client = None
        if self.anthropic_available:
            try:
                self.client = anthropic.Anthropic(api_key=self.api_key)
                logger.info("✅ Anthropic client initialized")
            except Exception as e:
                logger.warning(f"⚠️ Anthropic init failed: {e}")
                self.anthropic_available = False

    def check_plagiarism_local(self, text1: str, text2: str) -> Tuple[float, dict]:
        if not text1 or not text2:
            return 0.0, {"error": "Both texts required"}
        
        t1 = re.sub(r"[^\w\s]", "", text1.lower())
        t2 = re.sub(r"[^\w\s]", "", text2.lower())
        overall_similarity = SequenceMatcher(None, t1, t2).ratio()
        
        # Simple scoring logic
        score = overall_similarity * 100
        
        details = {
            "overall_similarity": score,
            "word_count_1": len(text1.split()),
            "word_count_2": len(text2.split()),
        }
        return score, details

    def get_plagiarism_report(self, text1: str, text2: str) -> Tuple[str, float, str]:
        score, details = self.check_plagiarism_local(text1, text2)
        status = "High Similarity" if score > 50 else "Low Similarity"
        report = f"Plagiarism Score: {score:.1f}% - {status}"
        return report, score, status

    def get_plagiarism_statistics(self, text1: str, text2: str) -> Dict:
        score, details = self.check_plagiarism_local(text1, text2)
        return {"overall_score": score, "details": details}

    def find_similar_sentences(self, text1: str, text2: str, threshold: float = 0.7) -> List[dict]:
        sentences1 = [s.strip() for s in re.split(r"[.!?]+", text1) if s.strip()]
        sentences2 = [s.strip() for s in re.split(r"[.!?]+", text2) if s.strip()]
        pairs = []
        for i, s1 in enumerate(sentences1):
            for j, s2 in enumerate(sentences2):
                sim = SequenceMatcher(None, s1.lower(), s2.lower()).ratio()
                if sim >= threshold:
                    pairs.append({
                        "text1": s1, "text2": s2, "similarity": sim * 100
                    })
        return pairs

    # Dummy method placeholders if your route calls them but they were missing in the provided code
    def detect_plagiarism_patterns(self, text: str):
        # Placeholder logic
        return {"general": ["No specific patterns implemented yet"]}

    def calculate_plagiarism_risk(self, text: str, reference: str):
        score, _ = self.check_plagiarism_local(text, reference)
        risk_level = "HIGH" if score > 50 else "LOW"
        return score, {"risk_level": risk_level, "risk_factors": [], "detected_patterns": []}
    
    def get_detection_report(self, text: str, reference: str):
        score, details = self.calculate_plagiarism_risk(text, reference)
        return "Analysis Report", score, details

    def remove(self, text: str, rewrite_mode: str = "balanced", reference_text: str = "") -> Tuple[str, Dict]:
        # Simple fallback or Anthropic logic here
        if self.client:
            prompt = f"Rewrite this text ({rewrite_mode}) to avoid plagiarism:\n{text}"
            try:
                resp = self.client.messages.create(
                    model="claude-sonnet-4-5-20250929",
                    max_tokens=1000,
                    messages=[{"role": "user", "content": prompt}]
                )
                return resp.content[0].text, {"status": "anthropic_success"}
            except Exception:
                pass
        return text, {"status": "fallback_no_change"}
